import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

def format_sheet(ws, sdg_number):
    """Formata uma aba do Excel conforme o template desejado."""
    # Definir estilos
    header_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    metric_font = Font(bold=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    
    # Título do SDG
    sdg_titles = {
        1: "SDG1: No Poverty",
        2: "SDG2: Zero Hunger",
        3: "SDG3: Good Health and Well-being",
        # Adicione outros SDGs conforme necessário
        17: "SDG17: Partnerships for the Goals"
    }
    
    title = sdg_titles.get(sdg_number, f"SDG{sdg_number}")
    ws.cell(row=1, column=1, value=title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    
    title_cell = ws.cell(row=1, column=1)
    title_cell.fill = header_fill
    title_cell.font = Font(color="FFFFFF", bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Cabeçalhos da tabela
    headers = [
        "Type",
        "Metric and indicator reference",
        "Metric / Indicator", 
        "Value (for continuous data)",
        "Yes/No",
        "Evidence",
        "Public (Yes/No)"
    ]
    
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Ajustar largura das colunas
    col_widths = [12, 12, 50, 15, 12, 12, 12]
    for idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    
    # Altura das linhas
    for row in range(1, ws.max_row + 1):
        ws.row_dimensions[row].height = 25
    
    # Aplicar bordas
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border

def process_csv_files(csv_files):
    """Processa os arquivos CSV e organiza por SDG, Metric e Indicator."""
    sdg_structure = {}
    
    for csv_file in csv_files:
        try:
            # Extrair informações do nome do arquivo
            parts = csv_file.stem.split('_')
            if len(parts) < 2:
                continue
                
            # Exemplo: indicator_1-2-1 → sdg=1, metric=2, indicator=1
            numbers = parts[1].split('-')
            sdg = int(numbers[0])
            metric = int(numbers[1]) if len(numbers) > 1 else None
            indicator = int(numbers[2]) if len(numbers) > 2 else None
            
            # Ler o CSV
            df = pd.read_csv(csv_file)
            
            # Adicionar ao dicionário de estrutura
            if sdg not in sdg_structure:
                sdg_structure[sdg] = {}
            
            if metric:
                if metric not in sdg_structure[sdg]:
                    sdg_structure[sdg][metric] = {}
                
                if indicator:
                    sdg_structure[sdg][metric][indicator] = df
                else:
                    sdg_structure[sdg][metric]['main'] = df
            else:
                sdg_structure[sdg]['main'] = df
                
        except Exception as e:
            print(f"Erro ao processar {csv_file.name}: {e}")
    
    return sdg_structure

def write_to_excel(sdg_structure, excel_path):
    """Escreve a estrutura organizada no arquivo Excel."""
    with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
        for sdg, metrics in sdg_structure.items():
            # Criar uma aba para cada SDG
            sheet_name = f"sdg{sdg}"
            ws = writer.book.create_sheet(title=sheet_name)
            
            # Inicializar contador de linhas
            current_row = 3  # Começa após título e cabeçalhos
            
            # Formatar a aba
            format_sheet(ws, sdg)
            
            # Processar métricas e indicadores
            for metric_num, metric_data in metrics.items():
                if metric_num == 'main':
                    # Dados principais do SDG (sem métrica específica)
                    for _, row in metric_data.iterrows():
                        for col, value in enumerate(row, start=1):
                            ws.cell(row=current_row, column=col, value=value)
                        current_row += 1
                    continue
                
                # Adicionar métrica
                metric_ref = f"{sdg}.{metric_num}"
                ws.cell(row=current_row, column=1, value="Metric")
                ws.cell(row=current_row, column=2, value=metric_ref)
                
                # Mesclar células da descrição da métrica
                metric_desc = metric_data.get('main', pd.DataFrame()).iloc[0, 0] if 'main' in metric_data else ""
                ws.cell(row=current_row, column=3, value=metric_desc)
                ws.merge_cells(start_row=current_row, start_column=3, end_row=current_row, end_column=7)
                
                # Formatar linha da métrica
                for col in range(1, 8):
                    ws.cell(row=current_row, column=col).font = Font(bold=True)
                
                current_row += 1
                
                # Adicionar indicadores
                for indicator_num, indicator_df in metric_data.items():
                    if indicator_num == 'main':
                        continue
                    
                    # Referência do indicador
                    indicator_ref = f"{sdg}.{metric_num}.{indicator_num}"
                    
                    # Adicionar cada linha do DataFrame do indicador
                    for _, row in indicator_df.iterrows():
                        ws.cell(row=current_row, column=1, value="Indicator")
                        ws.cell(row=current_row, column=2, value=indicator_ref)
                        
                        # Copiar os outros valores
                        for col in range(3, min(8, len(row) + 2)):
                            ws.cell(row=current_row, column=col, value=row[col-2])
                        
                        current_row += 1

def generate_formatted_excel():
    # Criar pasta se não existir
    excel_dir = Path('scripts/excel')
    excel_dir.mkdir(parents=True, exist_ok=True)
    
    # Verificar se a pasta data existe
    data_dir = Path('data')
    if not data_dir.exists():
        print(f"Erro: A pasta {data_dir} não existe!")
        return
    
    # Listar CSVs
    csv_files = list(data_dir.glob('indicator_*.csv'))
    
    if not csv_files:
        print(f"Nenhum arquivo CSV encontrado em {data_dir}")
        return
    
    print(f"Encontrados {len(csv_files)} arquivos CSV em {data_dir}")
    
    # Processar arquivos e criar estrutura
    sdg_structure = process_csv_files(csv_files)
    
    # Criar arquivo Excel
    excel_path = excel_dir / 'sdg-structured-data.xlsx'
    write_to_excel(sdg_structure, excel_path)
    
    print(f"\nExcel formatado gerado com sucesso: {excel_path}")
    print(f"Abas criadas: {list(sdg_structure.keys())}")

if __name__ == "__main__":
    generate_formatted_excel()
